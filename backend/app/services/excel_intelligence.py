"""M03 — Excel Intelligence & Column Mapping.

Pure functions over bytes/openpyxl objects — no FastAPI, no database — so header
detection and mapping suggestion are unit-testable without an HTTP layer or storage.

Every suggestion here is exactly that: a suggestion. Nothing in this module writes to
the database; the router only persists a mapping once an auditor calls the confirm
endpoint (spec section 10: 'Never automatically approve a fuzzy AI match as a final
audit decision' applies just as much to column mapping as material matching).
"""

import difflib
import io
from dataclasses import dataclass, field

import openpyxl

from app.models.enums import DocumentType

# Canonical target fields per document type, each with header aliases used for
# suggestion matching. Order doesn't matter; matching is alias-driven, not positional.
CANONICAL_FIELDS: dict[DocumentType, dict[str, list[str]]] = {
    DocumentType.IMPORT_MIS: {
        "be_number": ["be no", "b/e no", "bill of entry no", "be number", "boe no"],
        "be_date": ["be date", "b/e date", "bill of entry date", "boe date"],
        "import_date": ["import date", "date of import", "arrival date"],
        "lc_number": ["lc no", "l/c no", "lc number"],
        "invoice_number": ["invoice no", "invoice number", "commercial invoice no"],
        "item_description": ["description", "item description", "goods description", "particulars"],
        "hs_code": ["hs code", "h.s. code", "hscode", "tariff heading"],
        "declared_quantity": ["quantity", "qty", "declared quantity", "import quantity"],
        "declared_unit": ["unit", "uom", "unit of measure"],
        "original_currency": ["currency", "curr"],
        "original_currency_value": ["value", "invoice value", "fob value", "cif value"],
    },
    DocumentType.ENTITLEMENT_SHEET: {
        "group_name": ["group", "group name", "category"],
        "group_code": ["group code", "category code"],
        "sequence_no": ["sl", "sl no", "serial", "sl.no", "si no"],
        "material_name": ["material", "material name", "item name", "raw material name"],
        "hs_code": ["hs code", "h.s. code", "hscode"],
        "entitlement_unit": ["unit", "uom", "unit of measure"],
        "annual_entitlement_qty": ["annual entitlement", "approved qty", "annual approved quantity"],
        "enhanced_entitlement_qty": ["enhanced entitlement", "enhanced qty", "additional entitlement"],
    },
    DocumentType.ENHANCED_ENTITLEMENT: {
        "material_name": ["material", "material name"],
        "hs_code": ["hs code", "h.s. code"],
        "enhanced_entitlement_qty": ["enhanced entitlement", "enhanced qty"],
    },
}

MAX_HEADER_SCAN_ROWS = 15


@dataclass
class HeaderDetectionResult:
    sheet_name: str
    header_row_index: int  # 0-based, within the scanned rows
    headers: list[str]


@dataclass
class MappingSuggestion:
    source_header: str
    suggested_target_field: str | None
    confidence: float


@dataclass
class ParsedRow:
    row_number: int  # 1-based, matches the spreadsheet's own row numbering
    values: dict[str, str | int | float | None] = field(default_factory=dict)


def list_sheets(content: bytes) -> list[str]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def detect_header_row(content: bytes, sheet_name: str) -> HeaderDetectionResult:
    """Heuristic: within the first MAX_HEADER_SCAN_ROWS rows, the header row is the one
    with the most non-empty, non-numeric string cells — spreadsheet headers are text,
    while rows above them (titles/company names) tend to have few populated cells and
    rows below them (data) tend to include numbers."""
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        best_index = 0
        best_score = -1
        best_headers: list[str] = []
        for idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS, values_only=True)):
            text_cells = [str(cell).strip() for cell in row if isinstance(cell, str) and str(cell).strip()]
            score = len(text_cells)
            if score > best_score:
                best_score = score
                best_index = idx
                best_headers = [str(cell).strip() if cell is not None else "" for cell in row]
        # Trim trailing empty header columns.
        while best_headers and best_headers[-1] == "":
            best_headers.pop()
        return HeaderDetectionResult(sheet_name=sheet_name, header_row_index=best_index, headers=best_headers)
    finally:
        workbook.close()


def _normalize(text: str) -> str:
    return " ".join(text.lower().replace(".", " ").replace("_", " ").replace("/", " ").split())


def suggest_column_mapping(headers: list[str], document_type: DocumentType) -> list[MappingSuggestion]:
    field_aliases = CANONICAL_FIELDS.get(document_type, {})
    normalized_alias_lookup: list[tuple[str, str]] = [
        (_normalize(alias), target_field) for target_field, aliases in field_aliases.items() for alias in aliases
    ]

    suggestions: list[MappingSuggestion] = []
    for header in headers:
        if not header:
            suggestions.append(MappingSuggestion(source_header=header, suggested_target_field=None, confidence=0.0))
            continue
        normalized_header = _normalize(header)

        exact = next((target for alias, target in normalized_alias_lookup if alias == normalized_header), None)
        if exact:
            suggestions.append(MappingSuggestion(source_header=header, suggested_target_field=exact, confidence=1.0))
            continue

        candidates = [alias for alias, _ in normalized_alias_lookup]
        close = difflib.get_close_matches(normalized_header, candidates, n=1, cutoff=0.6)
        if close:
            matched_alias = close[0]
            target = next(target for alias, target in normalized_alias_lookup if alias == matched_alias)
            confidence = difflib.SequenceMatcher(None, normalized_header, matched_alias).ratio()
            suggestions.append(
                MappingSuggestion(source_header=header, suggested_target_field=target, confidence=round(confidence, 4))
            )
        else:
            suggestions.append(MappingSuggestion(source_header=header, suggested_target_field=None, confidence=0.0))
    return suggestions


def read_rows(
    content: bytes,
    sheet_name: str,
    header_row_index: int,
    headers: list[str],
    limit: int | None = None,
) -> list[ParsedRow]:
    """Rows below the header, keyed by the raw header text (Layer A: unmapped, untouched)."""
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows: list[ParsedRow] = []
        data_start = header_row_index + 2  # openpyxl min_row is 1-based; data begins the row after the header
        for excel_row_number, row in enumerate(sheet.iter_rows(min_row=data_start, values_only=True), start=data_start):
            if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
                continue
            values = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            rows.append(ParsedRow(row_number=excel_row_number, values=values))
            if limit is not None and len(rows) >= limit:
                break
        return rows
    finally:
        workbook.close()
