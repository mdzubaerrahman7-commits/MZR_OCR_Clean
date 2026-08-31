"""M10 / OUTPUT 01-07 — Reporting & Findings Engine (spec sections 16-17).

Every generator here is a pure function: plain dataclasses in, an openpyxl
`Workbook` out. No database, no FastAPI — the router (reports.py) queries the
current audit state, builds these dataclasses, and calls the matching generator.
That keeps 'never sort the final entitlement-based report alphabetically' and
'show group subtotals and grand totals' (spec section 17) things a unit test can
assert on the workbook structure directly.

Group/item order in every output is exactly the caller-supplied sequence — nothing
in this module re-sorts.
"""

from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True)
SUBTOTAL_FONT = Font(bold=True, italic=True)
GRAND_TOTAL_FONT = Font(bold=True)


@dataclass(frozen=True)
class ReportEntitlementItem:
    entitlement_item_id: str
    group_sequence_no: int
    group_name: str
    item_sequence_no: int
    material_name: str
    hs_code: str
    entitlement_unit: str
    annual_entitlement_qty: Decimal
    enhanced_entitlement_qty: Decimal
    total_entitlement_qty: Decimal


@dataclass(frozen=True)
class ReportTransactionLine:
    entitlement_item_id: str | None
    material_name: str
    hs_code: str
    be_number: str
    be_date: date
    declared_quantity: Decimal
    declared_unit: str
    entitlement_quantity: Decimal | None
    entitlement_unit: str | None
    kg_quantity: Decimal | None
    original_currency: str | None
    original_currency_value: Decimal | None
    usd_value: Decimal | None
    classification: str


@dataclass(frozen=True)
class ReportFindingLine:
    finding_code: str
    issue_type: str
    evidence: str
    be_number: str | None
    material_name: str | None
    quantity: Decimal | None
    unit: str | None
    value: Decimal | None
    demand_amount: Decimal | None
    review_status: str


@dataclass(frozen=True)
class ExceptionSummary:
    non_entitled_count: int
    review_required_match_count: int
    unknown_classification_count: int
    partial_excess_count: int
    full_excess_count: int
    conversion_incomplete_count: int


def _write_header(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT


def _autosize(sheet: Worksheet, headers: list[str]) -> None:
    # Column letters are computed from the column index directly rather than read off
    # a row-1 cell: row 1 holds the merged report title, so every column past the
    # first is a MergedCell there, which doesn't expose `.column_letter`.
    for col, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = max(12, len(header) + 2)


def _report_title(sheet: Worksheet, title: str, span: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    cell = sheet.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal="center")


# --- OUTPUT 01 — Raw Material Detailed Import Statement ---------------------------

OUTPUT_01_HEADERS = [
    "Group", "SL", "Material Name", "HS Code", "B/E No", "B/E Date",
    "Declared Qty", "Declared Unit", "Entitlement Qty", "Entitlement Unit",
    "KG Qty", "Currency", "Currency Value", "USD Value",
]


def generate_output_01(
    items: list[ReportEntitlementItem], lines_by_item: dict[str, list[ReportTransactionLine]], audit_title: str
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output 01 - Raw Material Detail"
    _report_title(sheet, f"{audit_title} — Raw Material Detailed Import Statement", len(OUTPUT_01_HEADERS))
    row = 3
    _write_header(sheet, row, OUTPUT_01_HEADERS)
    row += 1

    current_group = None
    group_ent_total = group_kg_total = group_usd_total = Decimal("0")
    for item in items:
        if current_group is not None and item.group_name != current_group:
            _write_output_01_subtotal(sheet, row, current_group, group_ent_total, group_kg_total, group_usd_total)
            row += 1
            group_ent_total = group_kg_total = group_usd_total = Decimal("0")
        current_group = item.group_name

        for line in lines_by_item.get(item.entitlement_item_id, []):
            sheet.append(
                [
                    item.group_name, item.item_sequence_no, item.material_name, item.hs_code,
                    line.be_number, line.be_date.isoformat(),
                    float(line.declared_quantity), line.declared_unit,
                    float(line.entitlement_quantity) if line.entitlement_quantity is not None else None,
                    line.entitlement_unit,
                    float(line.kg_quantity) if line.kg_quantity is not None else None,
                    line.original_currency,
                    float(line.original_currency_value) if line.original_currency_value is not None else None,
                    float(line.usd_value) if line.usd_value is not None else None,
                ]
            )
            row += 1
            group_ent_total += line.entitlement_quantity or Decimal("0")
            group_kg_total += line.kg_quantity or Decimal("0")
            group_usd_total += line.usd_value or Decimal("0")

    if current_group is not None:
        _write_output_01_subtotal(sheet, row, current_group, group_ent_total, group_kg_total, group_usd_total)

    _autosize(sheet, OUTPUT_01_HEADERS)
    return workbook


def _write_output_01_subtotal(sheet: Worksheet, row: int, group_name: str, ent_total: Decimal, kg_total: Decimal, usd_total: Decimal) -> None:
    sheet.cell(row=row, column=1, value=f"{group_name} — Subtotal").font = SUBTOTAL_FONT
    sheet.cell(row=row, column=9, value=float(ent_total)).font = SUBTOTAL_FONT
    sheet.cell(row=row, column=11, value=float(kg_total)).font = SUBTOTAL_FONT
    sheet.cell(row=row, column=14, value=float(usd_total)).font = SUBTOTAL_FONT


# --- OUTPUT 02 — Imported Raw Material Summary -------------------------------------

OUTPUT_02_HEADERS = [
    "Group", "SL", "Material Name", "HS Code", "Entitlement Unit",
    "Annual Entitlement", "Enhanced Entitlement", "Total Entitlement",
    "Consumed (Entitlement Unit)", "Consumed (KG)", "Consumed (USD)", "Balance",
]


def generate_output_02(
    items: list[ReportEntitlementItem], lines_by_item: dict[str, list[ReportTransactionLine]], audit_title: str
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output 02 - RM Summary"
    _report_title(sheet, f"{audit_title} — Imported Raw Material Summary", len(OUTPUT_02_HEADERS))
    row = 3
    _write_header(sheet, row, OUTPUT_02_HEADERS)
    row += 1

    current_group = None
    group_totals = [Decimal("0")] * 5  # annual, enhanced, total, consumed_ent, consumed_kg (usd tracked separately below)
    group_usd = Decimal("0")
    grand_totals = [Decimal("0")] * 5
    grand_usd = Decimal("0")

    def flush_group_subtotal(group_name: str) -> None:
        nonlocal row
        sheet.cell(row=row, column=1, value=f"{group_name} — Subtotal").font = SUBTOTAL_FONT
        for offset, value in enumerate(group_totals):
            sheet.cell(row=row, column=6 + offset, value=float(value)).font = SUBTOTAL_FONT
        sheet.cell(row=row, column=11, value=float(group_usd)).font = SUBTOTAL_FONT
        row += 1

    for item in items:
        if current_group is not None and item.group_name != current_group:
            flush_group_subtotal(current_group)
            group_totals = [Decimal("0")] * 5
            group_usd = Decimal("0")
        current_group = item.group_name

        lines = lines_by_item.get(item.entitlement_item_id, [])
        consumed_ent = sum((l.entitlement_quantity or Decimal("0") for l in lines), Decimal("0"))
        consumed_kg = sum((l.kg_quantity or Decimal("0") for l in lines), Decimal("0"))
        consumed_usd = sum((l.usd_value or Decimal("0") for l in lines), Decimal("0"))
        balance = item.total_entitlement_qty - consumed_ent

        sheet.append(
            [
                item.group_name, item.item_sequence_no, item.material_name, item.hs_code, item.entitlement_unit,
                float(item.annual_entitlement_qty), float(item.enhanced_entitlement_qty), float(item.total_entitlement_qty),
                float(consumed_ent), float(consumed_kg), float(consumed_usd), float(balance),
            ]
        )
        row += 1

        item_values = [item.annual_entitlement_qty, item.enhanced_entitlement_qty, item.total_entitlement_qty, consumed_ent, consumed_kg]
        group_totals = [a + b for a, b in zip(group_totals, item_values, strict=True)]
        group_usd += consumed_usd
        grand_totals = [a + b for a, b in zip(grand_totals, item_values, strict=True)]
        grand_usd += consumed_usd

    if current_group is not None:
        flush_group_subtotal(current_group)

    sheet.cell(row=row, column=1, value="GRAND TOTAL").font = GRAND_TOTAL_FONT
    for offset, value in enumerate(grand_totals):
        sheet.cell(row=row, column=6 + offset, value=float(value)).font = GRAND_TOTAL_FONT
    sheet.cell(row=row, column=11, value=float(grand_usd)).font = GRAND_TOTAL_FONT

    _autosize(sheet, OUTPUT_02_HEADERS)
    return workbook


# --- OUTPUT 03 — Machinery & Sample Import Statement -------------------------------

OUTPUT_03_HEADERS = ["B/E No", "B/E Date", "Description", "HS Code", "Quantity", "Unit", "Currency", "Currency Value", "USD Value"]


def generate_output_03(machinery_lines: list[ReportTransactionLine], sample_lines: list[ReportTransactionLine], audit_title: str) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_title, lines in (("Machinery", machinery_lines), ("Sample", sample_lines)):
        sheet = workbook.create_sheet(title=f"Output 03 - {sheet_title}")
        _report_title(sheet, f"{audit_title} — {sheet_title} Import Statement", len(OUTPUT_03_HEADERS))
        _write_header(sheet, 3, OUTPUT_03_HEADERS)
        for line in lines:
            sheet.append(
                [
                    line.be_number, line.be_date.isoformat(), line.material_name, line.hs_code,
                    float(line.declared_quantity), line.declared_unit, line.original_currency,
                    float(line.original_currency_value) if line.original_currency_value is not None else None,
                    float(line.usd_value) if line.usd_value is not None else None,
                ]
            )
        _autosize(sheet, OUTPUT_03_HEADERS)
    return workbook


# --- OUTPUT 04 — Reconciliation & Exception Dashboard ------------------------------


def generate_output_04(summary: ExceptionSummary, audit_title: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output 04 - Exceptions"
    _report_title(sheet, f"{audit_title} — Import Audit Reconciliation & Exception Dashboard", 2)
    rows = [
        ("Non-entitled imports (no approved entitlement mapping)", summary.non_entitled_count),
        ("Matches awaiting auditor review", summary.review_required_match_count),
        ("Transactions with Unknown classification", summary.unknown_classification_count),
        ("Partial excess transactions", summary.partial_excess_count),
        ("Full excess transactions", summary.full_excess_count),
        ("Approved matches with incomplete quantity/currency conversion", summary.conversion_incomplete_count),
    ]
    row = 3
    _write_header(sheet, row, ["Audit Control", "Count"])
    for label, count in rows:
        row += 1
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=count)
    _autosize(sheet, ["Audit Control", "Count"])
    return workbook


# --- OUTPUT 05 / 06 — B/E-wise Duty Assessment (non-entitled / excess) -------------

ASSESSMENT_HEADERS = ["Finding Code", "B/E No", "Material", "Quantity", "Unit", "Value (USD)", "Demand Amount", "Review Status", "Evidence"]


def _generate_assessment_output(findings: list[ReportFindingLine], sheet_title: str, report_title: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    _report_title(sheet, report_title, len(ASSESSMENT_HEADERS))
    row = 3
    _write_header(sheet, row, ASSESSMENT_HEADERS)
    total_demand = Decimal("0")
    for finding in findings:
        sheet.append(
            [
                finding.finding_code, finding.be_number, finding.material_name,
                float(finding.quantity) if finding.quantity is not None else None,
                finding.unit,
                float(finding.value) if finding.value is not None else None,
                float(finding.demand_amount) if finding.demand_amount is not None else None,
                finding.review_status, finding.evidence,
            ]
        )
        total_demand += finding.demand_amount or Decimal("0")
        row += 1
    row += 1
    sheet.cell(row=row, column=1, value="TOTAL DEMAND").font = GRAND_TOTAL_FONT
    sheet.cell(row=row, column=7, value=float(total_demand)).font = GRAND_TOTAL_FONT
    _autosize(sheet, ASSESSMENT_HEADERS)
    return workbook


def generate_output_05(non_entitled_findings: list[ReportFindingLine], audit_title: str) -> Workbook:
    return _generate_assessment_output(
        non_entitled_findings, "Output 05 - Non-Entitled Duty", f"{audit_title} — Non-Entitled Import B/E-wise Duty Assessment"
    )


def generate_output_06(excess_findings: list[ReportFindingLine], audit_title: str) -> Workbook:
    return _generate_assessment_output(
        excess_findings, "Output 06 - Excess Duty", f"{audit_title} — Excess Import B/E-wise Duty Assessment"
    )


# --- OUTPUT 07 — Import Audit Findings Register ------------------------------------

OUTPUT_07_HEADERS = [
    "Finding Code", "Issue Type", "B/E No", "Material", "Quantity", "Unit",
    "Value (USD)", "Demand Amount", "Review Status", "Evidence",
]


def generate_output_07(findings: list[ReportFindingLine], audit_title: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output 07 - Findings Register"
    _report_title(sheet, f"{audit_title} — Import Audit Findings Register", len(OUTPUT_07_HEADERS))
    row = 3
    _write_header(sheet, row, OUTPUT_07_HEADERS)
    for finding in findings:
        sheet.append(
            [
                finding.finding_code, finding.issue_type, finding.be_number, finding.material_name,
                float(finding.quantity) if finding.quantity is not None else None,
                finding.unit,
                float(finding.value) if finding.value is not None else None,
                float(finding.demand_amount) if finding.demand_amount is not None else None,
                finding.review_status, finding.evidence,
            ]
        )
    _autosize(sheet, OUTPUT_07_HEADERS)
    return workbook
